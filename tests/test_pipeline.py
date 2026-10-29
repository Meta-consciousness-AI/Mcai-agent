# -*- coding: utf-8 -*-
"""端到端冒烟测试：用合成语料验证抽取 → 构建 → 站点 → 数据集全链路。"""
import json
import sys
import shutil
import unittest
from pathlib import Path

from docx import Document

ROOT = Path(__file__).resolve().parent.parent
sys.path.insert(0, str(ROOT))

from mcai import dataset, site  # noqa: E402
from mcai.corpus import load_corpus  # noqa: E402
from mcai.extract import build  # noqa: E402

PROJECT = {
    "name": "测试引擎", "short_name": "Test", "brand": "元意识",
    "org": "测试机构", "site_url": "http://127.0.0.1:8787",
    "license": "CC BY-NC-SA 4.0", "version": "0.1.0",
}


def make_docx(path: Path) -> None:
    doc = Document()
    doc.add_heading("测试术语表", level=1)
    doc.add_paragraph("元意识体系")
    doc.add_paragraph("元意识：元意识是宇宙的本源，不可描述。大致可以用无形无相、不生不灭来辅助认知。")
    doc.add_paragraph("物质：一直处在稳定状态下稳定结构的元态能量。")
    doc.add_paragraph("意识：一直处在离散状态下离散运行的元态能量。")
    doc.save(path)


class PipelineTest(unittest.TestCase):
    def setUp(self) -> None:
        # 不用 tempfile：Windows 沙盒下 tmp 目录无法写入，改用固定工作目录
        self.dir = Path(__file__).parent / "_work"
        if self.dir.exists():
            shutil.rmtree(self.dir, ignore_errors=True)
        self.dir.mkdir(parents=True)
        make_docx(self.dir / "terms.docx")
        (self.dir / "config").mkdir()
        (self.dir / "config" / "corpus.json").write_text(json.dumps({
            "project": PROJECT,
            "corpus": [{"id": "terms", "path": str(self.dir / "terms.docx"),
                        "title": "测试术语表", "kind": "glossary", "priority": 100}],
            "keywords": ["元意识"],
            "term_blacklist": [],
            "slugs": {"元意识": "meta-consciousness"},
        }, ensure_ascii=False), encoding="utf-8")

    def tearDown(self) -> None:
        shutil.rmtree(self.dir, ignore_errors=True)

    def test_extract(self) -> None:
        corpus = load_corpus(self.dir / "config" / "corpus.json")
        kb = build(corpus, {"categories": {"元意识体系": ["元意识", "物质", "意识"]}})
        terms = {c.term for c in kb.concepts}
        self.assertIn("元意识", terms)
        self.assertIn("物质", terms)
        concept = kb.by_term["元意识"]
        self.assertEqual(concept.slug, "meta-consciousness")
        self.assertEqual(concept.category, "元意识体系")
        self.assertEqual(concept.status, "ok")
        self.assertTrue(any(f.question.startswith("什么是元意识") for f in kb.faqs))

    def test_site_and_dataset(self) -> None:
        corpus = load_corpus(self.dir / "config" / "corpus.json")
        kb = build(corpus, {"categories": {"元意识体系": ["元意识", "物质", "意识"]}})
        site.build_site(kb, self.dir / "dist")
        dataset.export(kb, self.dir / "dist" / "dataset")
        page = (self.dir / "dist" / "concepts" / "meta-consciousness.html").read_text(encoding="utf-8")
        self.assertIn("DefinedTerm", page)
        self.assertIn("元意识是宇宙的本源", page)
        jsonl = (self.dir / "dist" / "dataset" / "concepts.jsonl").read_text(encoding="utf-8")
        self.assertIn('"term": "元意识"', jsonl)
        self.assertTrue((self.dir / "dist" / "llms.txt").exists())
        self.assertTrue((self.dir / "dist" / "sitemap.xml").exists())


if __name__ == "__main__":
    unittest.main()

