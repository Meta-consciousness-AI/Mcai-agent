# -*- coding: utf-8 -*-
"""自动驾驶测试：指纹检测、自动重建、变更日志、状态文件。"""
import json
import shutil
import sys
import unittest
from pathlib import Path

from docx import Document

ROOT = Path(__file__).resolve().parent.parent
sys.path.insert(0, str(ROOT))

from mcai import workspace  # noqa: E402
from mcai.autopilot import Autopilot  # noqa: E402


class AutopilotTest(unittest.TestCase):
    def setUp(self) -> None:
        self.work = Path(__file__).parent / "_work_auto"
        shutil.rmtree(self.work, ignore_errors=True)
        workspace.init_workspace(self.work)
        self._write_docx(["元意识：元意识是宇宙的本源，不可描述。", "物质：一直处在稳定结构的元态能量。"])
        cfg_path = self.work / "config" / "corpus.json"
        cfg = json.loads(cfg_path.read_text(encoding="utf-8"))
        cfg["corpus"] = [{"id": "t", "path": "corpus/t.docx", "title": "测试词表",
                          "kind": "glossary", "priority": 100}]
        cfg_path.write_text(json.dumps(cfg, ensure_ascii=False), encoding="utf-8")
        self.pilot = Autopilot(self.work, ensure_server=False)

    def tearDown(self) -> None:
        shutil.rmtree(self.work, ignore_errors=True)

    def _write_docx(self, lines: list[str]) -> None:
        doc = Document()
        doc.add_heading("测试词表", level=1)
        for line in lines:
            doc.add_paragraph(line)
        doc.save(self.work / "corpus" / "t.docx")

    def test_cycle_detects_change(self) -> None:
        first = self.pilot.run_once()
        self.assertTrue(first["rebuilt"])
        self.assertTrue(first["status"]["ok"])
        self.assertTrue((self.work / "data" / "status.json").exists())
        self.assertTrue((self.work / "dist" / "status.json").exists())
        # 重建后应自动生成一份快照
        self.assertTrue(list((self.work / "backups").glob("snapshot-*.zip")))

        second = self.pilot.run_once()
        self.assertFalse(second["rebuilt"], "语料未变时不应重复构建")

        self._write_docx([
            "元意识：元意识是宇宙的本源，不可描述。",
            "物质：一直处在稳定结构的元态能量。",
            "意识：一直处在离散运行的元态能量。",
        ])
        third = self.pilot.run_once()
        self.assertTrue(third["rebuilt"])
        self.assertEqual(third["diff"]["added"], ["意识"])

        changelog = (self.work / "data" / "changelog.jsonl").read_text(encoding="utf-8").strip().splitlines()
        self.assertTrue(changelog)
        entry = json.loads(changelog[-1])
        self.assertIn("意识", entry["added"])

    def test_status_fields(self) -> None:
        self.pilot.run_once()
        status = json.loads((self.work / "data" / "status.json").read_text(encoding="utf-8"))
        for key in ("ok", "checked_at", "concepts", "server_ok", "rebuilt", "error"):
            self.assertIn(key, status)
        self.assertGreaterEqual(status["concepts"], 2)


if __name__ == "__main__":
    unittest.main()
