"""把项目内的 Markdown 文档转换为 docx（标题、段落、列表、表格、粗体）。

用法：
    python tools/md_to_docx.py 输入.md 输出.docx ["可选标题"]
"""
from __future__ import annotations

import re
import sys
from pathlib import Path

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.shared import Cm, Pt


def _set_cn_font(run, name="宋体", size=None, bold=None):
    run.font.name = name
    run.font.element.rPr.rFonts.set(qn("w:eastAsia"), name)
    if size is not None:
        run.font.size = Pt(size)
    if bold is not None:
        run.bold = bold


def _add_inline(paragraph, text: str, size=10.5):
    """处理 **粗体** 与 `代码`。"""
    for part in re.split(r"(\*\*.+?\*\*|`.+?`)", text):
        if not part:
            continue
        if part.startswith("**") and part.endswith("**"):
            run = paragraph.add_run(part[2:-2])
            _set_cn_font(run, size=size, bold=True)
        elif part.startswith("`") and part.endswith("`"):
            run = paragraph.add_run(part[1:-1])
            run.font.name = "Consolas"
            run.font.size = Pt(size - 1)
        else:
            run = paragraph.add_run(part)
            _set_cn_font(run, size=size)


def _is_table_row(line: str) -> bool:
    return line.strip().startswith("|") and line.strip().endswith("|")


def _cells(line: str) -> list[str]:
    return [c.strip() for c in line.strip().strip("|").split("|")]


def convert(md_path: Path, docx_path: Path, title: str = "") -> None:
    doc = Document()
    section = doc.sections[0]
    section.top_margin = Cm(2.54)
    section.bottom_margin = Cm(2.54)
    section.left_margin = Cm(2.5)
    section.right_margin = Cm(2.5)

    normal = doc.styles["Normal"]
    normal.font.name = "宋体"
    normal.font.size = Pt(10.5)
    normal.element.rPr.rFonts.set(qn("w:eastAsia"), "宋体")
    normal.paragraph_format.line_spacing = 1.5
    normal.paragraph_format.space_after = Pt(6)

    lines = md_path.read_text(encoding="utf-8").splitlines()
    i = 0
    first_heading = True
    while i < len(lines):
        raw = lines[i]
        line = raw.strip()
        if not line:
            i += 1
            continue

        if line == "---":
            i += 1
            continue

        heading = re.match(r"^(#{1,4})\s+(.*)$", line)
        if heading:
            level = len(heading.group(1))
            text = heading.group(2).strip()
            if level == 1 and first_heading and title:
                text = title
            p = doc.add_paragraph()
            p.paragraph_format.space_before = Pt(12)
            p.paragraph_format.space_after = Pt(8)
            if level == 1:
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                run = p.add_run(text)
                _set_cn_font(run, "黑体", 18, True)
                first_heading = False
            else:
                sizes = {2: 14, 3: 12, 4: 11}
                run = p.add_run(text)
                _set_cn_font(run, "黑体", sizes.get(level, 11), True)
            i += 1
            continue

        if _is_table_row(line):
            rows = []
            while i < len(lines) and _is_table_row(lines[i].strip()):
                row = _cells(lines[i])
                if not all(re.fullmatch(r":?-{2,}:?", c) for c in row):
                    rows.append(row)
                i += 1
            if rows:
                cols = max(len(r) for r in rows)
                table = doc.add_table(rows=len(rows), cols=cols)
                table.style = "Table Grid"
                table.alignment = WD_TABLE_ALIGNMENT.CENTER
                for r_idx, row in enumerate(rows):
                    for c_idx in range(cols):
                        cell = table.cell(r_idx, c_idx)
                        text = row[c_idx] if c_idx < len(row) else ""
                        cell.text = ""
                        p = cell.paragraphs[0]
                        _add_inline(p, text, size=9.5)
                        if r_idx == 0:
                            for run in p.runs:
                                run.bold = True
                doc.add_paragraph()
            continue

        bullet = re.match(r"^[-*]\s+(.*)$", line)
        if bullet:
            p = doc.add_paragraph(style="List Bullet")
            _add_inline(p, bullet.group(1))
            i += 1
            continue

        numbered = re.match(r"^\d+\.\s+(.*)$", line)
        if numbered:
            p = doc.add_paragraph(style="List Number")
            _add_inline(p, numbered.group(1))
            i += 1
            continue

        if line.startswith("```"):
            i += 1
            code_lines = []
            while i < len(lines) and not lines[i].strip().startswith("```"):
                code_lines.append(lines[i])
                i += 1
            i += 1
            for cl in code_lines:
                p = doc.add_paragraph()
                p.paragraph_format.left_indent = Cm(0.5)
                p.paragraph_format.line_spacing = 1.0
                run = p.add_run(cl)
                run.font.name = "Consolas"
                run.font.size = Pt(8.5)
            continue

        if line.startswith("> "):
            p = doc.add_paragraph()
            p.paragraph_format.left_indent = Cm(0.5)
            _add_inline(p, line[2:], size=10)
            i += 1
            continue

        # 普通段落：把连续的非特殊行合并
        buf = [line]
        i += 1
        while i < len(lines):
            nxt = lines[i].strip()
            if (not nxt or nxt == "---" or nxt.startswith("#") or _is_table_row(nxt)
                    or re.match(r"^[-*]\s+", nxt) or re.match(r"^\d+\.\s+", nxt)):
                break
            buf.append(nxt)
            i += 1
        p = doc.add_paragraph()
        p.paragraph_format.first_line_indent = Cm(0.74)
        _add_inline(p, " ".join(buf))

    doc.save(docx_path)


def main() -> int:
    if len(sys.argv) < 3:
        print("用法: python tools/md_to_docx.py 输入.md 输出.docx [标题]")
        return 2
    src, dst = Path(sys.argv[1]), Path(sys.argv[2])
    title = sys.argv[3] if len(sys.argv) > 3 else ""
    convert(src, dst, title)
    print(f"written: {dst}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
